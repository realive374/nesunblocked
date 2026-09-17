#!/usr/bin/env python3
"""Build 'Smile and Snarl: Crittrrific Switcheroo - Book Five' as a .docx."""
import sys, os, re, glob, json, copy
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.enum.section import WD_SECTION
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

BASE = os.path.dirname(os.path.abspath(__file__))

TITLE      = "Smile and Snarl"
SERIES     = "Crittrrific Switcheroo"
BOOKNUM    = "Book Five"
SUBTITLE   = "The Eight O'Clock Hour"
AUTHORLINE = ""

# ---------------------------------------------------------------- manuscript
def load_manuscript():
    blocks = []
    for path in sorted(glob.glob(os.path.join(BASE, "ms_*.txt"))):
        with open(path, encoding="utf-8") as fh:
            for raw in fh.read().split("\n"):
                line = raw.rstrip()
                if line.startswith("@PART|"):
                    _, label, name = line.split("|", 2)
                    blocks.append(("part", label, name))
                elif line.startswith("@CH|"):
                    _, num, name = line.split("|", 2)
                    blocks.append(("chapter", num, name))
                elif line.strip():
                    blocks.append(("para", line.strip(), None))
    return blocks

# ---------------------------------------------------------------- xml helpers
def field(par, instr):
    r = par.add_run()
    fc = OxmlElement('w:fldChar'); fc.set(qn('w:fldCharType'), 'begin'); r._r.append(fc)
    r2 = par.add_run()
    it = OxmlElement('w:instrText'); it.set(qn('xml:space'), 'preserve'); it.text = instr
    r2._r.append(it)
    r3 = par.add_run()
    fc2 = OxmlElement('w:fldChar'); fc2.set(qn('w:fldCharType'), 'end'); r3._r.append(fc2)
    return [r, r2, r3]

PPR_ORDER = ["pStyle","keepNext","keepLines","pageBreakBefore","framePr","widowControl",
  "numPr","suppressLineNumbers","pBdr","shd","tabs","suppressAutoHyphens","kinsoku",
  "wordWrap","overflowPunct","topLinePunct","autoSpaceDE","autoSpaceDN","bidi",
  "adjustRightInd","snapToGrid","spacing","ind","contextualSpacing","mirrorIndents",
  "suppressOverlap","jc","textDirection","textAlignment","textboxTightWrap","outlineLvl",
  "divId","cnfStyle","rPr","sectPr","pPrChange"]

RPR_ORDER = ["rStyle","rFonts","b","bCs","i","iCs","caps","smallCaps","strike","dstrike",
  "outline","shadow","emboss","imprint","noProof","snapToGrid","vanish","webHidden","color",
  "spacing","w","kern","position","sz","szCs","highlight","u","effect","bdr","shd","fitText",
  "vertAlign","rtl","cs","em","lang","eastAsianLayout","specVanish","oMath"]

def insert_ordered(parent, el, order):
    """Insert el into parent at the position required by the OOXML schema."""
    tag = el.tag.split('}')[-1]
    try:
        rank = order.index(tag)
    except ValueError:
        parent.append(el); return el
    for child in parent:
        ctag = child.tag.split('}')[-1]
        crank = order.index(ctag) if ctag in order else -1
        if crank > rank:
            child.addprevious(el)
            return el
    parent.append(el)
    return el

def set_tabs(par, positions):
    """positions: list of (inches, align, leader)"""
    pPr = par._p.get_or_add_pPr()
    tabs = OxmlElement('w:tabs')
    for pos, align, leader in positions:
        t = OxmlElement('w:tab')
        t.set(qn('w:val'), align)
        t.set(qn('w:pos'), str(int(pos * 1440)))
        if leader:
            t.set(qn('w:leader'), leader)
        tabs.append(t)
    insert_ordered(pPr, tabs, PPR_ORDER)

def keep_with_next(par):
    pPr = par._p.get_or_add_pPr()
    insert_ordered(pPr, OxmlElement('w:keepNext'), PPR_ORDER)

def no_widows(par):
    pPr = par._p.get_or_add_pPr()
    w = OxmlElement('w:widowControl'); w.set(qn('w:val'), '1')
    insert_ordered(pPr, w, PPR_ORDER)

def char_spacing(run, twentieths):
    rPr = run._r.get_or_add_rPr()
    sp = OxmlElement('w:spacing'); sp.set(qn('w:val'), str(int(twentieths)))
    insert_ordered(rPr, sp, RPR_ORDER)

# ---------------------------------------------------------------- build
def build(out_path, P, toc_pages=None):
    doc = Document()

    sec = doc.sections[0]
    sec.page_width   = Inches(6.0)
    sec.page_height  = Inches(9.0)
    sec.top_margin    = Inches(P["top"])
    sec.bottom_margin = Inches(P["bottom"])
    sec.left_margin   = Inches(P["side"])
    sec.right_margin  = Inches(P["side"])
    sec.header_distance = Inches(0.4)
    sec.footer_distance = Inches(0.45)
    sec.different_first_page_header_footer = True

    # base style
    st = doc.styles["Normal"]
    st.font.name = P["font"]
    st.font.size = Pt(P["size"])
    st.element.rPr.rFonts.set(qn('w:eastAsia'), P["font"])
    pf = st.paragraph_format
    pf.line_spacing = P["lead"]
    pf.space_after  = Pt(0)
    pf.space_before = Pt(0)
    pf.alignment    = WD_ALIGN_PARAGRAPH.LEFT
    pf.first_line_indent = Inches(P["indent"])

    # footer page number
    footer = sec.footer
    fp = footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fp.paragraph_format.first_line_indent = Inches(0)
    field(fp, " PAGE ")
    for r in fp.runs:
        r.font.name = P["font"]; r.font.size = Pt(P["size"] - 3)
        r.font.color.rgb = RGBColor(0x44, 0x44, 0x44)

    # ---- helpers bound to this doc
    def para(text="", *, size=None, bold=False, italic=False, align=None,
             indent=None, before=0, after=0, lead=None, caps=False,
             spacing=None, color=None, font=None, keep=False):
        p = doc.add_paragraph()
        pfm = p.paragraph_format
        pfm.alignment = align if align is not None else WD_ALIGN_PARAGRAPH.LEFT
        pfm.first_line_indent = Inches(indent if indent is not None else P["indent"])
        pfm.space_before = Pt(before)
        pfm.space_after  = Pt(after)
        pfm.line_spacing = lead if lead is not None else P["lead"]
        if keep: keep_with_next(p)
        no_widows(p)
        if text:
            r = p.add_run(text.upper() if caps else text)
            r.font.name = font or P["font"]
            r.font.size = Pt(size or P["size"])
            r.bold = bold; r.italic = italic
            if color is not None: r.font.color.rgb = color
            if spacing is not None:
                char_spacing(r, spacing * 20)
        return p

    def pagebreak():
        p = doc.add_paragraph()
        p.paragraph_format.first_line_indent = Inches(0)
        p.add_run().add_break(WD_BREAK.PAGE)
        return p

    def blank(n=1, size=None):
        for _ in range(n):
            para("", size=size or P["size"], align=WD_ALIGN_PARAGRAPH.CENTER, indent=0)

    GREY = RGBColor(0x55, 0x55, 0x55)

    # ============================================================ FRONT MATTER
    # 1. half title
    blank(6)
    para(TITLE, size=P["size"]+8, align=WD_ALIGN_PARAGRAPH.CENTER, indent=0,
         caps=True, spacing=2.4, lead=1.0)
    blank(1)
    para(SERIES, size=P["size"]-1, align=WD_ALIGN_PARAGRAPH.CENTER, indent=0,
         italic=True, color=GREY, lead=1.0)
    pagebreak()

    # 2. also by
    blank(5)
    para("Also in this series", size=P["size"]-1, align=WD_ALIGN_PARAGRAPH.CENTER,
         indent=0, caps=True, spacing=1.6, color=GREY, lead=1.0)
    blank(2)
    for t in ["Book One", "Book Two", "Book Three", "Book Four"]:
        para(t, size=P["size"], align=WD_ALIGN_PARAGRAPH.CENTER, indent=0,
             italic=True, lead=1.0, after=8)
    blank(2)
    para("Book Five", size=P["size"], align=WD_ALIGN_PARAGRAPH.CENTER, indent=0,
         bold=True, lead=1.0)
    pagebreak()

    # 3. title page
    blank(5)
    para(TITLE, size=P["size"]+12, align=WD_ALIGN_PARAGRAPH.CENTER, indent=0,
         caps=True, spacing=3.0, lead=1.0)
    blank(1)
    para(SERIES, size=P["size"]+1, align=WD_ALIGN_PARAGRAPH.CENTER, indent=0,
         italic=True, lead=1.0)
    blank(3)
    para(BOOKNUM, size=P["size"]-1, align=WD_ALIGN_PARAGRAPH.CENTER, indent=0,
         caps=True, spacing=2.2, color=GREY, lead=1.0)
    blank(1)
    para(SUBTITLE, size=P["size"]+4, align=WD_ALIGN_PARAGRAPH.CENTER, indent=0,
         bold=True, lead=1.0)
    pagebreak()

    # 4. copyright
    blank(10)
    for line in [
        "This is a work of fan fiction, written for private, non-commercial enjoyment.",
        "",
        "Playcare, Ms Delight, Bobby BearHug, CatNap, DogDay, Hoppy Hopscotch,",
        "Bubba Bubbaphant, CraftyCorn, KickinChicken and PickyPiggy are the",
        "creations of their respective owners and are used here affectionately,",
        "without claim of ownership and without authorisation.",
        "",
        "\"SML\" refers to the real YouTube channel of that name. Its depiction",
        "here is invented for the purposes of this story. Nothing in this book is",
        "intended as a statement of fact about that channel or the people who make it.",
        "",
        "All other characters and events are invented.",
        "",
        "No part of this story is about a real school, a real teacher, or a real child.",
    ]:
        para(line, size=P["size"]-3, align=WD_ALIGN_PARAGRAPH.CENTER, indent=0,
             color=GREY, lead=1.2, after=1)
    pagebreak()

    # 5. dedication
    blank(9)
    para("For everybody who got asked the hard question", size=P["size"]+1,
         align=WD_ALIGN_PARAGRAPH.CENTER, indent=0, italic=True, lead=1.3)
    para("and then got left alone with it.", size=P["size"]+1,
         align=WD_ALIGN_PARAGRAPH.CENTER, indent=0, italic=True, lead=1.3)
    pagebreak()

    # 6. the story so far
    blank(3)
    para("The Story So Far", size=P["size"]+5, align=WD_ALIGN_PARAGRAPH.CENTER,
         indent=0, bold=True, lead=1.0)
    blank(2)
    for line in [
        "Playcare is a bright, enormous, underground sort of town where eleven hundred "
        "children live, go to school, and are looked after by eight Smiling Critters.",
        "DogDay leads. Hoppy Hopscotch runs. CraftyCorn paints. Bubba Bubbaphant thinks "
        "things through. KickinChicken skates. PickyPiggy negotiates about toast. "
        "CatNap gets everybody to sleep.",
        "And Bobby BearHug hugs people. That is not a small job. She has a certificate.",
        "In the four books before this one, Playcare survived the Great Switcheroo, the "
        "second Great Switcheroo, and the one everybody agreed not to call the third "
        "Great Switcheroo even though that is obviously what it was.",
        "This book has no switcheroo in it.",
        "This one is about a word.",
    ]:
        para(line, size=P["size"]-1, align=WD_ALIGN_PARAGRAPH.LEFT,
             indent=P["indent"], lead=P["lead"], after=4)
    pagebreak()

    # 7. TOC
    blocks = load_manuscript()
    blank(2)
    para("Contents", size=P["size"]+5, align=WD_ALIGN_PARAGRAPH.CENTER,
         indent=0, bold=True, lead=1.0)
    blank(2)
    text_w = 6.0 - 2 * P["side"]
    ch_i = 0
    for kind, a, b in blocks:
        if kind == "part":
            p = para("", indent=0, lead=1.0, before=9, after=4)
            r = p.add_run(f"{a}  ·  {b}")
            r.font.name = P["font"]; r.font.size = Pt(P["size"]-3)
            r.bold = True
            char_spacing(r, 24)
        elif kind == "chapter":
            pg = toc_pages[ch_i] if toc_pages and ch_i < len(toc_pages) else None
            ch_i += 1
            p = doc.add_paragraph()
            p.paragraph_format.first_line_indent = Inches(0)
            p.paragraph_format.left_indent = Inches(0.12)
            p.paragraph_format.space_after = Pt(2.0)
            p.paragraph_format.line_spacing = 1.08
            set_tabs(p, [(text_w - 0.12, 'right', 'dot')])
            r = p.add_run(f"{a}.  {b}")
            r.font.name = P["font"]; r.font.size = Pt(P["size"]-2.5)
            if pg:
                r2 = p.add_run("\t" + str(pg))
                r2.font.name = P["font"]; r2.font.size = Pt(P["size"]-2.5)
    pagebreak()

    # ============================================================ BODY
    chapter_marks = []   # (marker text) to locate in pdf
    first = True
    for kind, a, b in blocks:
        if kind == "part":
            if not first: pagebreak()
            first = False
            blank(7)
            para(a, size=P["size"]-1, align=WD_ALIGN_PARAGRAPH.CENTER, indent=0,
                 caps=True, spacing=3.2, color=GREY, lead=1.0)
            blank(1)
            para(b, size=P["size"]+6, align=WD_ALIGN_PARAGRAPH.CENTER, indent=0,
                 bold=True, lead=1.1)
            pagebreak()
            blank(14)
            pagebreak()
        elif kind == "chapter":
            if not first: pagebreak()
            first = False
            blank(3)
            para(f"Chapter {a}", size=P["size"]-3, align=WD_ALIGN_PARAGRAPH.CENTER,
                 indent=0, caps=True, spacing=2.6, color=GREY, lead=1.0, keep=True)
            blank(1)
            p = para(b, size=P["size"]+4, align=WD_ALIGN_PARAGRAPH.CENTER, indent=0,
                     bold=True, lead=1.15, keep=True)
            blank(1)
            chapter_marks.append((a, b))
            pend = True
        else:
            txt = a
            if txt == "THE END":
                blank(2)
                para("THE END", size=P["size"]+1, align=WD_ALIGN_PARAGRAPH.CENTER,
                     indent=0, caps=True, spacing=3.0, bold=True, lead=1.2, before=10)
            else:
                para(txt, after=P["pspace"])

    # ============================================================ BACK MATTER
    pagebreak()
    blank(3)
    para("Questions for the Eight O'Clock Hour", size=P["size"]+3,
         align=WD_ALIGN_PARAGRAPH.CENTER, indent=0, bold=True, lead=1.15)
    blank(1)
    para("Fenn puts three of these up on the board every Tuesday. "
         "Nobody has to answer any of them. “I don't know yet” is a real answer "
         "and always has been.",
         size=P["size"]-2, align=WD_ALIGN_PARAGRAPH.CENTER, indent=0,
         italic=True, color=GREY, lead=1.3, after=10)
    blank(1)
    QS = [
     "Ms Delight never raises her voice once in this whole book. Does that make what she did better, worse, or neither?",
     "Gus says he likes SML because “it's never worse than you thought.” What do you watch for that reason?",
     "The Arrangement on the fridge was a good rule and it lasted thirty-one hours. Was it a failure?",
     "Jacob puts the 2019 “media literacy” line into evidence even though it hurts his own case. Why?",
     "Six hundred and eighty-two children said the grown-ups were right to worry. Would you have been one of them?",
     "Bobby refuses to run a petition because “a petition tells people what they think.” Is she right?",
     "Mr Voss tells the Panel he can't stand behind his own client's sentence. What did that cost him?",
     "Jacob says he doesn't want to win a case people can call fixed. Is that wisdom or is it pride?",
     "The council repealed the ban because of CatNap's tally book, not because of the hearing. Why that and not everything else?",
     "The first Eight O'Clock night was a shambles. What would you have done on the second one?",
     "Marlow stopped watching because of something his brother said at dinner, not because of a law. What does that tell you?",
     "Ms Delight's mother wrote, “I don't know if that was right.” Was it?",
     "Bobby says the good bit was being asked and the bad bit was being left alone afterwards. Which happens to you more?",
     "The narrator never tells you the word. Did you want to know? Do you still?",
    ]
    for i, q in enumerate(QS, 1):
        p = doc.add_paragraph()
        p.paragraph_format.first_line_indent = Inches(-0.28)
        p.paragraph_format.left_indent = Inches(0.28)
        p.paragraph_format.space_after = Pt(7)
        p.paragraph_format.line_spacing = P["lead"] * 0.92
        p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
        no_widows(p)
        r = p.add_run(f"{i}. {q}")
        r.font.name = P["font"]; r.font.size = Pt(P["size"]-1)

    pagebreak()
    blank(3)
    para("A Note on Book Five", size=P["size"]+3, align=WD_ALIGN_PARAGRAPH.CENTER,
         indent=0, bold=True, lead=1.15)
    blank(2)
    for line in [
        "The first four books in this series were switcheroo books. Somebody woke up in "
        "somebody else's body, everybody learned something, and by the last chapter they "
        "were all back where they started.",
        "This one does the switcheroo differently, and if you have got this far you have "
        "probably already spotted it.",
        "Ms Delight starts this book as the person asking the hard question and ends it as "
        "the person being asked. Jacob starts it as the kid who writes things down and ends "
        "it as the one everybody quotes. Bobby BearHug starts it by being handled kindly in "
        "a small warm office, and ends it on a low wall outside the rec hall, handling "
        "somebody else kindly, with the exact same words.",
        "That is the switcheroo. Nobody changed bodies. They changed chairs.",
        "It is the only one in the series that nobody undid at the end, and that is on "
        "purpose, because it is the only one that was not an accident. Everybody in this "
        "book chose their swap, one hard conversation at a time, mostly at night, mostly "
        "while tired.",
        "That is how the real ones go.",
    ]:
        para(line, size=P["size"]-1, lead=P["lead"], after=5)

    doc.save(out_path)
    return chapter_marks

# ---------------------------------------------------------------- measure
def page_count_and_starts(docx_path, marks):
    import subprocess
    outdir = os.path.dirname(docx_path)
    pdf = docx_path[:-5] + ".pdf"
    txt = docx_path[:-5] + ".txt"
    for f in (pdf, txt):
        if os.path.exists(f): os.remove(f)
    subprocess.run(["soffice", "--headless", "--norestore", "--convert-to", "pdf",
                    "--outdir", outdir, docx_path], capture_output=True, timeout=900)
    if not os.path.exists(pdf):
        return None, None
    info = subprocess.run(["pdfinfo", pdf], capture_output=True, text=True).stdout
    n = None
    for line in info.splitlines():
        if line.lower().startswith("pages:"):
            n = int(line.split(":")[1].strip())
    subprocess.run(["pdftotext", "-layout", pdf, txt], capture_output=True, timeout=300)
    raw_pages = open(txt, encoding="utf-8", errors="replace").read().split("\f")
    starts = []
    for num, title in marks:
        pat = re.compile(r"^[ \t]*CHAPTER[ \t]+" + re.escape(str(num)) + r"[ \t]*$",
                         re.MULTILINE)
        found = None
        for i, pg in enumerate(raw_pages):
            if pat.search(pg):
                found = i + 1
                break
        starts.append(found)
    return n, starts

if __name__ == "__main__":
    P = json.loads(sys.argv[1]) if len(sys.argv) > 1 else {}
    D = dict(font="Times New Roman", size=12.5, lead=1.55, indent=0.26,
             top=0.7, bottom=0.65, side=0.72, pspace=0.0)
    D.update(P)
    out = os.path.join(BASE, "book5.docx")
    toc = None
    tocfile = os.path.join(BASE, "toc.json")
    if "--with-toc" in sys.argv and os.path.exists(tocfile):
        toc = json.load(open(tocfile))
    marks = build(out, D, toc_pages=toc)
    n, starts = page_count_and_starts(out, marks)
    json.dump(starts, open(os.path.join(BASE, "starts.json"), "w"))
    print(json.dumps({"pages": n, "missing_marks": sum(1 for s in starts if not s)}))
